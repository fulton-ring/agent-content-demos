import marimo

__generated_with = "0.25.1"
app = marimo.App(width="medium")

with app.setup:
    import io
    import json
    import os
    import re
    import time
    from datetime import UTC, datetime
    from pathlib import Path

    import geopandas as gpd
    import marimo as mo
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    import openpyxl
    import pandas as pd
    import requests
    import shapely
    from pypdf import PdfReader

    HERE = Path(__file__).parent
    DATA = Path(os.environ.get("SUPPLY_CHAIN_DATA_DIR", HERE / "data"))
    OUT = Path(os.environ.get("SUPPLY_CHAIN_OUTPUT_DIR", HERE / "output")) / "citywide"
    DATA.mkdir(parents=True, exist_ok=True)
    OUT.mkdir(parents=True, exist_ok=True)
    NY_FEET = "EPSG:2263"
    BOROUGHS = {"MN": "Manhattan", "BX": "Bronx", "BK": "Brooklyn", "QN": "Queens", "SI": "Staten Island"}
    RUN_STARTED = time.perf_counter()


@app.cell(hide_code=True)
def _():
    mo.md(r"""
    # Citywide: which substation serves every tax lot in New York City?

    **The question.** The [Brooklyn notebook](supply_chain_risk.py) showed that FEMA's supply chain mapping steps
    work for power in one borough: every lot traced to a Con Edison network or feeder, then to its area substation,
    transmission substation and plants, with a public source for each link. Does that hold for all five boroughs?

    Three things change at city scale:

    1. **Network to area substation.** Con Edison's 2012 network ranking, filed with the PSC, pairs all of the
       city's underground networks with their area substations. We parse it by code, compare it with the networks
       on Con Edison's map today, and record what moved since 2012.
    2. **Area substation to transmission substation.** Brooklyn needed a handful of filings. The city needs one
       link for every area substation, from Con Edison's and NYISO's public documents.
    3. **The places Con Edison's networks don't reach.** Most of Staten Island and eastern Queens are fed by
       overhead feeders at 33, 27, 13 and 4 kV, and the Rockaways are served by LIPA through PSEG Long Island,
       not by Con Edison at all.

    We measure what the issue asks for: facilities and links by type and borough, documented versus inferred;
    coverage against the counts Con Edison publishes per borough; and the share of the city's 857,000 tax lots
    matched to a network or feeder.
    """)
    return


@app.cell
def _():
    def get(url: str, params: dict | None = None, tries: int = 4) -> requests.Response:
        """GET with a fixed number of retries."""
        last = "no attempt"
        for attempt in range(tries):
            try:
                response = requests.get(
                    url,
                    params=params,
                    timeout=(20, 180),
                    headers={"User-Agent": "Mozilla/5.0 (nyc-supply-chain-risk spike)"},
                )
            except requests.RequestException as error:
                last = repr(error)
            else:
                if response.status_code == 200:
                    return response
                last = f"HTTP {response.status_code}"
            time.sleep(min(2**attempt, 15))
        raise RuntimeError(f"{url} failed after {tries} tries: {last}")

    def arcgis(layer: str, params: dict, page: int = 2000, max_pages: int = 200) -> gpd.GeoDataFrame:
        """All features from an ArcGIS REST layer query, paged, as WGS84 GeoJSON."""
        features: list[dict] = []
        for number in range(max_pages):
            query = {"where": "1=1", "outSR": 4326, "f": "geojson"} | params
            query |= {"resultOffset": number * page, "resultRecordCount": page}
            batch = get(f"{layer}/query", query).json()
            if "features" not in batch:
                raise RuntimeError(f"{layer} returned no features: {str(batch)[:300]}")
            features += batch["features"]
            if len(batch["features"]) < page:
                return gpd.GeoDataFrame.from_features(features, crs=4326)
        raise RuntimeError(f"{layer} has more than {max_pages * page} features")

    def cached(name: str, load) -> gpd.GeoDataFrame:
        """Run load() once and keep the result in DATA/name (GeoParquet) for later runs."""
        path = DATA / name
        if not path.exists():
            load().to_parquet(path)
        return gpd.read_parquet(path)

    return arcgis, cached, get


@app.cell(hide_code=True)
def _():
    mo.md(r"""
    ## 1. Sources, and checking each passage

    As in the Brooklyn notebook, every link that comes from a document carries the exact passage we rely on, and
    the notebook downloads the document and checks the passage is in it (ignoring case, spaces and punctuation).
    """)
    return


@app.cell
def _():
    _psc = "https://documents.dps.ny.gov/public/Common/ViewDoc.aspx?DocRefId="
    _nyiso = "https://www.nyiso.com/documents/20142"
    documents = {
        "psc-network-ranking-2012": (
            "Con Edison, 2012 Network Ranking, Exhibit IIP-8, filed with the NY PSC",
            f"{_psc}%7BF62626F2-2E5B-47D3-88FB-EC5AC20C085D%7D",
        ),
        "ju-reliability-2024": (
            "Joint Utilities of NY, 2024 Annual Report on Electric Service and Power Quality (Con Edison section)",
            "https://jointutilitiesofny.org/sites/default/files/2024%20Annual%20Report%20on%20Electric%20Service"
            "%20and%20Power%20Quality.pdf",
        ),
        "bqdm-petition-2014": (
            "Con Edison, Brooklyn/Queens Demand Management petition, NY PSC Case 14-E-0302, July 2014",
            f"{_psc}%7BFA6E8548-E790-4E6A-8BF2-61DDF62EAB4E%7D",
        ),
        "bceh-dps12-2023": (
            "Con Edison, response DPS-12 on the Brooklyn Clean Energy Hub, NY PSC Case 20-E-0197, February 2023",
            f"{_psc}%7BC0E87187-0000-CF35-9148-3E240230B194%7D",
        ),
        "nyiso-queue": (
            "NYISO Interconnection Queue (current file; withdrawn and in-service sheets)",
            f"{_nyiso}/1407078/NYISO-Interconnection-Queue.xlsx",
        ),
        "nyiso-bps": (
            "NYISO, The New York State Bulk Power Transmission Facilities",
            f"{_nyiso}/1407249/existing_bps_transmission_facilities.pdf/aecc0414-397b-4095-9d21-634eb9674639",
        ),
        "coned-ltp-2025": (
            "Con Edison, Local Transmission Plan update to NYISO, 21 January 2025",
            f"{_nyiso}/49295323/CECONY's_LTP_Update_1_21_2025.pdf/abf6cfb4-10e6-eee4-3988-0a434f5a1dcb",
        ),
        "nyiso-gold-book-2026": (
            "NYISO, 2026 Load & Capacity Data (Gold Book)",
            f"{_nyiso}/2226333/2026-Gold-Book-Public.pdf",
        ),
    }
    return (documents,)


@app.cell
def _(documents, get):
    def normalise(text: str) -> str:
        return re.sub(r"[^a-z0-9]", "", text.lower())

    def raw_text(source: str) -> list[str] | None:
        """The document's text: one entry per spreadsheet row, or per PDF page, or one for a web page.
        None if it can't be downloaded from here. The Brooklyn notebook's copies in DATA are reused."""
        path = DATA / f"doc-{source}"
        if not path.exists():
            try:
                path.write_bytes(get(documents[source][1], tries=2).content)
            except RuntimeError as error:
                print(f"unreachable: {source}: {error}")
                return None
        raw = path.read_bytes()
        if raw.startswith(b"%PDF"):
            return [page.extract_text() or "" for page in PdfReader(io.BytesIO(raw)).pages]
        if raw.startswith(b"PK"):
            book = openpyxl.load_workbook(io.BytesIO(raw), read_only=True)
            rows = (row for sheet in book.worksheets for row in sheet.iter_rows(values_only=True))
            return [" ".join(str(c) for c in row if c is not None) for row in rows]
        return [re.sub(r"<[^>]+>", " ", raw.decode("utf-8", "replace"))]

    document_pages = {source: raw_text(source) for source in documents}
    return document_pages, normalise


@app.cell(hide_code=True)
def _():
    mo.md(r"""
    ## 2. Network to area substation

    **The 2012 table.** Con Edison's 2012 network ranking (PSC exhibit IIP-8) is one page with one row per network:
    rank, region, network ID, network name, area substation, and two reliability scores. We read it with a regular
    expression. Network and substation names run together in the text, so we split them at the longest network
    name that also appears on Con Edison's map today.
    """)
    return


@app.cell
def _(arcgis, cached):
    _coned = "https://services.arcgis.com/ciPnsNFi1JLWVjva/arcgis/rest/services"
    networks = cached(
        "coned-networks.parquet",
        lambda: arcgis(f"{_coned}/CECONY_NetworkLCSD_Prod/FeatureServer/0", {"outFields": "NETWORK,BOROUGH"}, 1000),
    )
    network_areas = cached(
        "coned-network-supplied-areas.parquet",
        lambda: arcgis(f"{_coned}/CECONY_NwkSuppliedArea_Prod/FeatureServer/0", {"outFields": "OBJECTID"}, 1000),
    )
    networks = networks[networks["BOROUGH"] != "Westchester"].rename(columns={"NETWORK": "network", "BOROUGH": "borough"})
    networks[["network", "borough"]].groupby("borough").size().rename("load areas on Con Edison's map")
    return network_areas, networks


@app.cell
def _(document_pages, networks):
    _text = document_pages["psc-network-ranking-2012"]
    assert _text is not None, "the 2012 network ranking is the base table; it must download"
    _row = re.compile(r"^(\d+) ([BMQX]) (\d+[BMQX]) (.+?) (\d\.\d{6}) (\d\.\d\d)\s*$")
    # Names on today's map, plus the one 2012 name that has since gone (Randalls Island, misspelled in the exhibit).
    _names = {re.sub(r" - (BK|QN)| SN\d", "", n) for n in networks["network"]} | {"Randalls Insland"}
    _parsed = []
    for _line in _text[0].splitlines():
        _m = _row.match(_line.strip())
        if not _m:
            continue
        _words = _m.group(4).split()
        _cut = max(k for k in range(1, len(_words)) if " ".join(_words[:k]) in _names)
        _parsed.append(
            {
                "network_id": _m.group(3),
                "network_2012": " ".join(_words[:_cut]),
                "substation_2012": " ".join(_words[_cut:]),
                "passage": _line.strip(),
            }
        )
    ranking_2012 = pd.DataFrame(_parsed)
    mo.md(
        f"Parsed **{len(ranking_2012)} networks** and **{ranking_2012['substation_2012'].nunique()} substation "
        "names** from the 2012 table."
    )
    return (ranking_2012,)


@app.cell
def _():
    def canonical_station(name: str) -> str:
        """One spelling per area substation: the 2012 table, the hosting capacity map and the filings differ."""
        name = re.sub(r"\s*\((Bk|Bx|Mhtn)\)", "", name).replace("Greeenwood", "Greenwood")
        name = name.title() if name.isupper() else name
        name = re.sub(r"(\d)(St|Nd|Rd|Th)\b", lambda m: m.group(1) + m.group(2).lower(), name)
        name = re.sub(r"\s+", " ", name).replace(" Substation", "")
        name = re.sub(r"No\.\s*(\d)", r"No. \1", name)
        name = re.sub(r"^E\. ", "East ", re.sub(r"^W\. ", "West ", name))
        name = re.sub(r"\bSt\.?$|\bSt\.? (?=No)", "Street ", name).strip()
        name = re.sub(r"\bAve\. A$", "Avenue A", name)
        return re.sub(r"\s+", " ", name.replace("Hell Gate", "Hellgate")).strip()

    return (canonical_station,)


@app.cell(hide_code=True)
def _():
    mo.md(r"""
    ## 5. Which network or feeder serves each lot, citywide

    The rule is the Brooklyn notebook's, run on all of the city's tax lots (MapPLUTO points from NYC Open Data):
    a lot inside both a network's load area and Con Edison's network-supplied area is served by that network;
    any other lot takes the area substation of the nearest overhead feeder segment within 500 feet; a lot near no
    segment falls back to the load area it sits in, marked inferred. Con Edison's hosting capacity layers carry
    every overhead segment in its territory with its area substation, so this step needs no new source outside
    the Rockaways.
    """)
    return


@app.cell
def _(arcgis, cached):
    _coned = "https://services.arcgis.com/ciPnsNFi1JLWVjva/arcgis/rest/services"
    _nyc_box = {
        "outFields": "FEEDER_ID,SUBSTATION",
        "geometry": "-74.26,40.49,-73.69,40.92",
        "geometryType": "esriGeometryEnvelope",
        "inSR": 4326,
        "spatialRel": "esriSpatialRelIntersects",
    }
    # Layer 0 holds three-phase overhead segments, layer 1 one- and two-phase ones.
    feeders = cached(
        "coned-overhead-segments-nyc.parquet",
        lambda: pd.concat(
            [arcgis(f"{_coned}/CECONY_NodalHCV_Prod/FeatureServer/{n}", _nyc_box) for n in (0, 1)],
            ignore_index=True,
        ),
    )
    feeders
    return (feeders,)


@app.cell
def _(cached, get):
    def pluto_nyc() -> gpd.GeoDataFrame:
        rows: list[dict] = []
        for page in range(40):
            batch = get(
                "https://data.cityofnewyork.us/resource/64uk-42ks.json",
                {
                    "$select": "bbl,borough,latitude,longitude",
                    "$where": "latitude IS NOT NULL",
                    "$order": "bbl",
                    "$limit": 50000,
                    "$offset": page * 50000,
                },
            ).json()
            rows += batch
            if len(batch) < 50000:
                lots = pd.DataFrame(rows)
                lots["bbl"] = lots["bbl"].astype(float).astype("int64").astype(str)
                points = gpd.points_from_xy(lots["longitude"].astype(float), lots["latitude"].astype(float))
                return gpd.GeoDataFrame(lots[["bbl", "borough"]], geometry=points, crs=4326)
        raise RuntimeError("PLUTO returned more than 2,000,000 lots")

    lots = cached("pluto-nyc-lots.parquet", pluto_nyc)
    lots["borough"] = lots["borough"].map(BOROUGHS)
    lots.groupby("borough").size().rename("tax lots")
    return (lots,)


if __name__ == "__main__":
    app.run()
